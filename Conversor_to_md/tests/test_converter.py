"""Testes com documentos sintéticos em memória; nenhum dado de paciente."""
from io import BytesIO
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from converter import ConversionError, convert_document


def docx_bytes(document):
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def pdf_bytes(text=True, blank=False, encrypted=False):
    """Constrói um PDF mínimo com fonte padrão e conteúdo controlado."""
    writer = PdfWriter()
    page = writer.add_blank_page(width=600, height=800)
    if text:
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
                                 NameObject('/Subtype'): NameObject('/Type1'),
                                 NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({
            NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
        content = DecodedStreamObject()
        content.set_data(b'BT /F1 12 Tf 50 750 Td (Caso ficticio: 10 mg) Tj ET')
        page[NameObject('/Contents')] = content
    if blank:
        writer.add_blank_page(width=600, height=800)
    if encrypted:
        writer.encrypt('senha-de-teste')
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


class ConverterTests(unittest.TestCase):
    def test_docx_structure_order_accents_and_table(self):
        document = Document()
        document.add_heading('Caso fictício', level=1)
        document.add_paragraph('Observação: 10 mg; sem interpretação.')
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = 'Campo'
        table.cell(0, 1).text = 'Valor'
        table.cell(1, 0).text = 'A|B'
        table.cell(1, 1).text = 'Linha 1\nLinha 2'
        document.add_paragraph('Fim')
        result = convert_document('../../caso.DOCX', docx_bytes(document))
        self.assertEqual(result.filename, 'caso.md')
        self.assertIn('# Caso fictício', result.markdown)
        self.assertIn('10 mg', result.markdown)
        self.assertIn('A\\|B', result.markdown)
        self.assertIn('<br>', result.markdown)
        self.assertLess(result.markdown.index('Campo'), result.markdown.index('Fim'))

    def test_docx_lists(self):
        document = Document()
        document.add_paragraph('Primeiro', style='List Bullet')
        document.add_paragraph('Segundo', style='List Number')
        result = convert_document('caso.docx', docx_bytes(document))
        self.assertIn('- Primeiro', result.markdown)
        self.assertIn('1. Segundo', result.markdown)

    def test_docx_empty(self):
        with self.assertRaisesRegex(ConversionError, 'não contém texto'):
            convert_document('vazio.docx', docx_bytes(Document()))

    def test_markdown_and_html_are_literal(self):
        document = Document()
        document.add_paragraph('<script> ![foto](https://exemplo.test/a)')
        text = convert_document('a.docx', docx_bytes(document)).markdown
        self.assertNotIn('<script>', text)
        self.assertIn('\\!\\[foto\\]', text)

    def test_pdf_text(self):
        result = convert_document('a.pdf', pdf_bytes())
        self.assertIn('## Página 1', result.markdown)
        self.assertIn('10 mg', result.markdown)

    def test_pdf_partial(self):
        result = convert_document('a.pdf', pdf_bytes(blank=True))
        self.assertTrue(any('Conversão parcial' in item for item in result.warnings))
        self.assertIn('## Página 2', result.markdown)

    def test_pdf_without_text(self):
        with self.assertRaisesRegex(ConversionError, 'OCR'):
            convert_document('scan.pdf', pdf_bytes(text=False))

    def test_encrypted_pdf(self):
        with self.assertRaisesRegex(ConversionError, 'senha'):
            convert_document('a.pdf', pdf_bytes(encrypted=True))

    def test_page_limit(self):
        with patch('converter.MAX_PAGES', 1):
            with self.assertRaisesRegex(ConversionError, 'páginas'):
                convert_document('a.pdf', pdf_bytes(blank=True))

    def test_invalid_inputs(self):
        cases = [('a.doc', b'x'), ('a.pdf', b''), ('a.pdf', b'fake'),
                 ('a.docx', b'fake'), ('a.pdf', b'%PDF-corrompido')]
        for filename, data in cases:
            with self.subTest(filename=filename, data=data):
                with self.assertRaises(ConversionError):
                    convert_document(filename, data)

    def test_size_limit(self):
        with patch('converter.MAX_BYTES', 2):
            with self.assertRaisesRegex(ConversionError, '10 MiB'):
                convert_document('a.pdf', b'123')

    def test_zip_not_docx(self):
        buffer = BytesIO()
        with ZipFile(buffer, 'w') as archive:
            archive.writestr('test.txt', 'texto')
        with self.assertRaisesRegex(ConversionError, 'DOCX válido'):
            convert_document('a.docx', buffer.getvalue())

    def test_zip_expansion_limit(self):
        document = Document()
        document.add_paragraph('Caso sintético')
        with patch('converter.MAX_UNCOMPRESSED_BYTES', 10):
            with self.assertRaisesRegex(ConversionError, 'descompactado'):
                convert_document('a.docx', docx_bytes(document))


if __name__ == '__main__':
    unittest.main()
